"""
excel_core.py
=============
Logika murni Python (tidak bergantung Streamlit) untuk:
  - Menyusun teks description sesuai aturan kategori (bensin/drink/parkir).
  - Dedupe transaksi parkir dari screenshot Flazz terhadap struk fisik.
  - Mengurutkan hasil secara kronologis.
  - Mengisi ke template Excel FORM_REIMBURSE.
"""

import io
import re
from calendar import monthrange
from datetime import date

import openpyxl
import pandas as pd
from dateutil import parser as dateparser

DATA_START_ROW = 12          # baris pertama tabel data di template
DATA_END_ROW_DEFAULT = 61    # baris terakhir yang sudah disiapkan template
TOTAL_ROW_DEFAULT = 62       # baris 'Total' di template


def parse_date_safe(raw):
    if not raw:
        return None
    try:
        return dateparser.parse(str(raw), dayfirst=False, yearfirst=True).date()
    except Exception:
        try:
            return dateparser.parse(str(raw), dayfirst=True).date()
        except Exception:
            return None


def parse_nominal(raw) -> float:
    try:
        return float(re.sub(r"[^\d.\-]", "", str(raw)) or 0)
    except ValueError:
        return 0.0


def build_description(item: dict) -> str:
    """Susun kolom description sesuai aturan bisnis per kategori (dipastikan di
    Python, tidak hanya mengandalkan kepatuhan model AI terhadap instruksi)."""
    itype = (item.get("type") or "").strip().lower()

    if itype == "bensin":
        fuel = (item.get("fuel_type") or "").strip()
        liters = item.get("liters")
        if not fuel:
            return "-"
        if "pertalite" in fuel.lower():
            return fuel
        if liters:
            try:
                liters_fmt = f"{float(liters):g}"
            except (TypeError, ValueError):
                liters_fmt = str(liters)
            return f"{fuel} {liters_fmt} Liter"
        return fuel

    if itype == "drink":
        drink = (item.get("drink_name") or "").strip()
        outlet = (item.get("outlet_name") or "").strip()
        parts = [p for p in [drink, outlet] if p]
        return " - ".join(parts) if parts else "-"

    if itype == "parkir":
        loc = (item.get("location_name") or "").strip()
        return loc if loc else "-"

    return "-"


def filter_and_dedupe_flazz(receipt_items: list, flazz_items: list) -> list:
    """Gabungkan transaksi parkir dari screenshot Flazz TANPA duplikat dengan struk fisik.

    Aturan:
    - Baris "Top Up" selalu di-skip, tidak pernah dimasukkan.
    - Baris "Parking" yang tanggal & nominalnya sama persis dengan salah satu
      struk parkir (foto fisik) yang sudah diekstrak -> di-skip, karena struk
      fisiknya sudah mewakili transaksi itu (hindari dobel input).
    - Baris "Parking" yang tidak match struk manapun tetap dimasukkan sebagai
      kategori parkir dengan description "-" (diisi manual oleh pengguna nanti,
      kecuali sudah diisi otomatis dari fallback GPS -- lihat app.py).
    """
    receipt_keys = set()
    for item in receipt_items:
        if (item.get("type") or "").strip().lower() == "parkir":
            d = parse_date_safe(item.get("date"))
            nominal = parse_nominal(item.get("nominal"))
            if d is not None:
                receipt_keys.add((d, round(nominal)))

    result = []
    for f in flazz_items:
        ftype = (f.get("type") or "").strip().lower()
        if "top" in ftype:
            continue  # Top Up tidak perlu dimasukkan
        if "park" not in ftype:
            continue  # hanya proses baris parkir dari screenshot, kategori lain diabaikan
        d = parse_date_safe(f.get("date"))
        nominal = parse_nominal(f.get("nominal"))
        if d is not None and (d, round(nominal)) in receipt_keys:
            continue  # sudah terwakili oleh struk parkir fisik -> skip, hindari duplikat
        result.append(
            {
                "date": f.get("date"),
                "type": "parkir",
                "nominal": f.get("nominal"),
                "location_name": "-",
            }
        )
    return result


def build_rows(extracted_items: list) -> pd.DataFrame:
    rows = []
    for item in extracted_items:
        d = parse_date_safe(item.get("date"))
        nominal = parse_nominal(item.get("nominal", 0))
        rows.append(
            {
                "date": d,
                "category": item.get("type", "-"),
                "description": build_description(item),
                "nominal": nominal,
            }
        )
    df = pd.DataFrame(rows)
    # struk tanpa tanggal terbaca ditaruh paling akhir, bukan hilang
    df["_sort_key"] = df["date"].apply(lambda d: d if d else date.max)
    df = df.sort_values(by="_sort_key", ascending=True).drop(columns=["_sort_key"]).reset_index(drop=True)
    return df


def fill_excel_template(df: pd.DataFrame, template_path: str, header_info: dict) -> bytes:
    wb = openpyxl.load_workbook(template_path)
    ws = wb["FORM"]

    if header_info.get("name"):
        ws["C6"] = header_info["name"]
    if header_info.get("department"):
        ws["C7"] = header_info["department"]
    if header_info.get("purpose"):
        ws["C8"] = header_info["purpose"]
    if header_info.get("bank_acc"):
        ws["C9"] = header_info["bank_acc"]

    valid_dates = [d for d in df["date"] if d]
    ref = min(valid_dates) if valid_dates else date.today()
    period_start = date(ref.year, ref.month, 1)
    period_end = date(ref.year, ref.month, monthrange(ref.year, ref.month)[1])
    ws["E6"].number_format = "d mmm yyyy"
    ws["E7"].number_format = "d mmm yyyy"
    ws["E6"] = period_start
    ws["E7"] = period_end

    n = len(df)
    available_rows = DATA_END_ROW_DEFAULT - DATA_START_ROW + 1  # 50 baris tersedia di template

    if n > available_rows:
        extra = n - available_rows
        ws.insert_rows(DATA_END_ROW_DEFAULT + 1, amount=extra)
        for i in range(extra):
            src_row = DATA_END_ROW_DEFAULT
            dst_row = DATA_END_ROW_DEFAULT + 1 + i
            for col in range(2, 6):  # B..E
                src_cell = ws.cell(row=src_row, column=col)
                dst_cell = ws.cell(row=dst_row, column=col)
                dst_cell.number_format = src_cell.number_format
                dst_cell.font = src_cell.font.copy()
                dst_cell.border = src_cell.border.copy()
                dst_cell.fill = src_cell.fill.copy()
                dst_cell.alignment = src_cell.alignment.copy()
        total_row = TOTAL_ROW_DEFAULT + extra
        data_end_row = DATA_END_ROW_DEFAULT + extra
    else:
        total_row = TOTAL_ROW_DEFAULT
        data_end_row = DATA_END_ROW_DEFAULT

    for i, row in df.iterrows():
        r = DATA_START_ROW + i
        ws.cell(row=r, column=2, value=row["date"])
        ws.cell(row=r, column=3, value=row["category"])
        ws.cell(row=r, column=4, value=row["description"])
        ws.cell(row=r, column=5, value=row["nominal"])

    ws.cell(row=total_row, column=5).value = f"=SUM(E{DATA_START_ROW}:E{data_end_row})"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
